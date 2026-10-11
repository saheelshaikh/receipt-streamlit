"""CPU-only OCR and PDF rendering. No Apple, Node or Codex dependencies."""
import re
from pathlib import Path

_engine = None

def engine():
    global _engine
    if _engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _engine = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2)
    return _engine

def render_source(source, directory):
    from PIL import Image, ImageOps
    paths = []
    if source.suffix.lower() == '.pdf':
        import pypdfium2 as pdfium
        document = pdfium.PdfDocument(str(source))
        try:
            if not 1 <= len(document) <= 30:
                raise ValueError('Use a PDF with 1 to 30 pages. Split larger scans into smaller files.')
            for index in range(len(document)):
                page = document[index]
                bitmap = None
                try:
                    width, height = page.get_size()
                    bitmap = page.render(scale=1800 / max(width, height))
                    target = directory / f'page-{index+1:04}.png'
                    bitmap.to_pil().save(target)
                    paths.append(target)
                finally:
                    if bitmap is not None: bitmap.close()
                    page.close()
        finally:
            document.close()
    else:
        with Image.open(source) as source_image:
            im = ImageOps.exif_transpose(source_image).convert('RGB')
            im.thumbnail((2200, 2200))
            target = directory / 'page-0001.png'
            im.save(target)
            paths.append(target)
    return paths

def ocr(paths):
    import numpy as np
    import zxingcpp
    from PIL import Image
    reader = engine()
    pages = []
    for path in paths:
        record = {'path': str(path), 'lines': [], 'qr': []}
        try:
            with Image.open(path) as source:
                im = source.convert('RGB')
                width, height = im.size
                pixels = np.asarray(im)
                record['qr'] = [b.text for b in zxingcpp.read_barcodes(im)
                                if b.format == zxingcpp.BarcodeFormat.QRCode]
                if not record['qr']:
                    # Re-sampling helps QR modules damaged by scanned PDF rasterization.
                    gray=im.convert('L')
                    for scale in (.8, .6, 1.2):
                        resized=gray.resize((max(1,int(width*scale)),max(1,int(height*scale))))
                        record['qr'].extend(b.text for b in zxingcpp.read_barcodes(resized)
                                            if b.format == zxingcpp.BarcodeFormat.QRCode)
                    record['qr']=list(dict.fromkeys(record['qr']))
                result, _ = reader(pixels)
                for box, text, score in result or []:
                    xs, ys = zip(*box)
                    x, y = max(0, min(xs)/width), max(0, min(ys)/height)
                    w, h = (max(xs)-min(xs))/width, (max(ys)-min(ys))/height
                    line = dict(text=text, confidence=float(score), x=x, y=y, w=w, h=h)
                    # Retain the whole detected line as crop evidence; no guessed character boxes.
                    line['numbers'] = [dict(text=m.group(), x=x, y=y, w=w, h=h)
                                       for m in re.finditer(r'-?\d+[.,]\d{2}(?!\d)', text)]
                    record['lines'].append(line)
        except Exception as error:
            record['error'] = str(error)
        pages.append(record)
    return pages

def export_xlsx(rows, path):
    """Standalone exporter for the user's portable Python tool."""
    from datetime import date
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import FormulaRule
    book = Workbook()
    sheet = book.active
    sheet.title = 'Receipts'
    sheet.sheet_view.showGridLines = False
    sheet.append(['Date', 'Receipt number', 'Amount (SAR)'])
    for row in rows:
        sheet.append([date.fromisoformat(row['date']) if row['date'] else None,
                      row['receipt_number'] or None,
                      float(row['amount']) if row['amount'] else None])
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='center')
            if cell.row == 1:
                cell.fill = PatternFill('solid', fgColor='203950')
                cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            elif cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F4F8')
        sheet.row_dimensions[cells[0].row].height = 25
    for i in range(2, len(rows)+2):
        sheet.cell(i,1).number_format = 'dd.mm.yyyy'
        sheet.cell(i,2).number_format = '@'
        sheet.cell(i,3).number_format = '#,##0.00'
    for col, width in [('A',18), ('B',32), ('C',20)]:
        sheet.column_dimensions[col].width = width
    if rows:
        sheet.conditional_formatting.add(f'A2:C{len(rows)+1}',
            FormulaRule(formula=['ISBLANK(A2)'],fill=PatternFill('solid',fgColor='FFF2CD')))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:C{len(rows)+1}'
    book.save(path)

def check_dependencies():
    import openpyxl, pypdfium2, zxingcpp, reportlab, PIL, onnxruntime
    from PIL import Image, ImageDraw
    import numpy as np
    # Load packaged models and run inference during setup, not the first upload.
    im=Image.new('RGB',(400,100),'white')
    ImageDraw.Draw(im).text((20,20),'53.30',fill='black')
    engine()(np.asarray(im))
    print('PDF, OCR models, QR decoder, Excel and PDF export dependencies are ready.')
